"""Document parsing, chunking, and vector similarity for the local RAG app."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from typing import Iterable

import numpy as np
from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class TextSection:
    text: str
    page: int | None = None


@dataclass(frozen=True)
class RetrievedChunk:
    document_id: str
    filename: str
    chunk_index: int
    text: str
    embedding: np.ndarray
    page: int | None = None


def safe_filename(filename: str) -> str:
    """Return a display-safe basename for either Windows or POSIX paths."""
    normalized = filename.replace("\\", "/").strip()
    return PurePosixPath(normalized).name


def extract_sections(filename: str, content: bytes) -> list[TextSection]:
    """Extract readable text from a PDF or DOCX byte stream."""
    extension = PurePosixPath(filename.lower()).suffix
    if extension == ".pdf":
        reader = PdfReader(BytesIO(content))
        return [
            TextSection(text=page.extract_text() or "", page=page_number)
            for page_number, page in enumerate(reader.pages, start=1)
        ]

    if extension == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        text = "\n".join(part for part in paragraphs + table_rows if part)
        return [TextSection(text=text)] if text else []

    raise ValueError("Only PDF and DOCX files are supported.")


def chunk_text(text: str, chunk_size: int = 180, overlap: int = 30) -> list[str]:
    """Split text into word-based chunks with a small, intentional overlap."""
    if chunk_size < 1:
        raise ValueError("chunk_size must be at least 1.")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap must be non-negative and smaller than chunk_size.")

    words = text.split()
    if not words:
        return []

    step = chunk_size - overlap
    chunks = []
    for start in range(0, len(words), step):
        chunks.append(" ".join(words[start : start + chunk_size]))
        if start + chunk_size >= len(words):
            break
    return chunks


def cosine_similarity(query: np.ndarray, vectors: np.ndarray) -> np.ndarray:
    """Compute cosine similarity between one query vector and many vectors."""
    query = np.asarray(query, dtype=np.float32).reshape(-1)
    vectors = np.asarray(vectors, dtype=np.float32)
    if vectors.ndim != 2 or vectors.shape[1] != query.shape[0]:
        raise ValueError("vectors must be a 2D array matching the query dimensions.")

    denominator = np.linalg.norm(vectors, axis=1) * np.linalg.norm(query)
    scores = np.zeros(vectors.shape[0], dtype=np.float32)
    valid = denominator > 0
    scores[valid] = (vectors[valid] @ query) / denominator[valid]
    return scores


def chunks_for_sections(
    sections: Iterable[TextSection], chunk_size: int, overlap: int
) -> list[tuple[int, str, int | None]]:
    """Build chunk text and source metadata while preserving page numbers."""
    chunks: list[tuple[int, str, int | None]] = []
    for section in sections:
        for text in chunk_text(section.text, chunk_size, overlap):
            chunks.append((len(chunks), text, section.page))
    return chunks